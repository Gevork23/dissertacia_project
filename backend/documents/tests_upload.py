import os
import shutil
import tempfile
from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import override_settings
from django.urls import reverse
from docx import Document as DocxDocument
from pypdf import PdfWriter
from rest_framework import status
from rest_framework.test import APITestCase

from .models import Document, DocumentVersion
from .text_processing import normalize_text, sha256_hex

TEST_MEDIA_ROOT = tempfile.mkdtemp()


@override_settings(MEDIA_ROOT=TEST_MEDIA_ROOT)
class DocumentUploadAPITests(APITestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass()
        shutil.rmtree(TEST_MEDIA_ROOT, ignore_errors=True)

    def make_file(self, name: str, content: str, content_type: str = "text/plain"):
        return SimpleUploadedFile(
            name=name,
            content=content.encode("utf-8"),
            content_type=content_type,
        )

    def make_docx_file(self, name: str = "reglament.docx") -> SimpleUploadedFile:
        buffer = BytesIO()
        document = DocxDocument()
        document.add_paragraph("Административный регламент")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Срок"
        table.rows[0].cells[1].text = "5 дней"
        document.save(buffer)
        return SimpleUploadedFile(
            name=name,
            content=buffer.getvalue(),
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
        )

    def make_pdf_file(
        self,
        name: str = "reglament.pdf",
        text: str = "Hello PDF",
    ) -> SimpleUploadedFile:
        safe_text = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
        stream = f"BT\n/F1 18 Tf\n50 100 Td\n({safe_text}) Tj\nET".encode("latin-1")
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            (
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] "
                b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
            ),
            b"<< /Length "
            + str(len(stream)).encode()
            + b" >>\nstream\n"
            + stream
            + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]

        pdf_bytes = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for index, obj in enumerate(objects, start=1):
            offsets.append(len(pdf_bytes))
            pdf_bytes.extend(f"{index} 0 obj\n".encode())
            pdf_bytes.extend(obj)
            pdf_bytes.extend(b"\nendobj\n")

        xref_offset = len(pdf_bytes)
        pdf_bytes.extend(f"xref\n0 {len(objects) + 1}\n".encode())
        pdf_bytes.extend(b"0000000000 65535 f \n")
        for offset in offsets[1:]:
            pdf_bytes.extend(f"{offset:010d} 00000 n \n".encode())
        pdf_bytes.extend(
            (
                f"trailer\n<< /Root 1 0 R /Size {len(objects) + 1} >>\n"
                f"startxref\n{xref_offset}\n%%EOF\n"
            ).encode()
        )

        return SimpleUploadedFile(
            name=name,
            content=bytes(pdf_bytes),
            content_type="application/pdf",
        )

    def test_create_document_returns_phase4_ready_fields(self):
        response = self.client.post(
            reverse("documents-list"),
            {
                "title": "Приказ о приёме документов",
                "description": "Карточка документа без версий",
            },
            format="json",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["title"], "Приказ о приёме документов")
        self.assertEqual(response.data["versions_count"], 0)
        self.assertIsNone(response.data["latest_version_number"])

    def test_upload_first_version_assigns_number_and_saves_file(self):
        document = Document.objects.create(title="Административный регламент")

        response = self.client.post(
            reverse("documents-versions", kwargs={"pk": document.id}),
            {
                "file": self.make_file("reglament_v1.txt", "Первая редакция"),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["document"], document.id)
        self.assertEqual(response.data["version_number"], 1)
        self.assertEqual(response.data["source_filename"], "reglament_v1.txt")
        self.assertEqual(
            response.data["file_size"], len("Первая редакция".encode("utf-8"))
        )
        self.assertEqual(response.data["content_type"], "text/plain")

        version = DocumentVersion.objects.get(pk=response.data["id"])
        self.assertTrue(os.path.exists(version.file.path))
        self.assertIn(f"document_{document.id}", version.file.name)
        self.assertIn("version_1", version.file.name)

    def test_upload_populates_extracted_text_normalized_text_and_content_hash_for_txt(
        self,
    ):
        document = Document.objects.create(title="TXT документ")
        raw_text = "Строка 1\r\n\r\nСтрока 2"

        response = self.client.post(
            reverse("documents-versions", kwargs={"pk": document.id}),
            {
                "file": self.make_file("law.txt", raw_text),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        version = DocumentVersion.objects.get(pk=response.data["id"])
        expected_extracted = "Строка 1\n\nСтрока 2"
        expected_normalized = normalize_text(expected_extracted)

        self.assertEqual(version.extracted_text, expected_extracted)
        self.assertEqual(version.normalized_text, expected_normalized)
        self.assertEqual(version.content_hash, sha256_hex(expected_normalized))

    def test_upload_populates_extracted_text_for_docx(self):
        document = Document.objects.create(title="DOCX документ")

        response = self.client.post(
            reverse("documents-versions", kwargs={"pk": document.id}),
            {
                "file": self.make_docx_file(),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        version = DocumentVersion.objects.get(pk=response.data["id"])
        self.assertIn("Административный регламент", version.extracted_text)
        self.assertIn("Срок | 5 дней", version.extracted_text)
        self.assertEqual(version.content_hash, sha256_hex(version.normalized_text))

    def test_upload_populates_extracted_text_for_pdf_without_ocr(self):
        document = Document.objects.create(title="PDF документ")

        response = self.client.post(
            reverse("documents-versions", kwargs={"pk": document.id}),
            {
                "file": self.make_pdf_file(text="Hello PDF"),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        version = DocumentVersion.objects.get(pk=response.data["id"])
        self.assertIn("Hello PDF", version.extracted_text)
        self.assertEqual(version.content_hash, sha256_hex(version.normalized_text))

    def test_upload_rejects_pdf_without_extractable_text(self):
        document = Document.objects.create(title="Скан PDF")
        buffer = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=144)
        writer.write(buffer)

        response = self.client.post(
            reverse("documents-versions", kwargs={"pk": document.id}),
            {
                "file": SimpleUploadedFile(
                    name="scan.pdf",
                    content=buffer.getvalue(),
                    content_type="application/pdf",
                ),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("OCR is not supported", response.data["file"][0])

    def test_version_text_endpoint_returns_processing_fields(self):
        document = Document.objects.create(title="Документ")
        version = DocumentVersion.objects.create(
            document=document,
            version_number=1,
            source_filename="v1.txt",
            file=self.make_file("v1.txt", "Версия 1"),
            file_size=len("Версия 1".encode("utf-8")),
            content_type="text/plain",
            extracted_text="Версия 1",
            normalized_text="Версия 1",
            content_hash=sha256_hex("Версия 1"),
        )

        response = self.client.get(reverse("versions-text", kwargs={"pk": version.id}))

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(response.data["extracted_text"], "Версия 1")
        self.assertEqual(response.data["normalized_text"], "Версия 1")
        self.assertEqual(response.data["content_hash"], sha256_hex("Версия 1"))

    def test_upload_next_version_auto_increments_number(self):
        document = Document.objects.create(title="Инструкция")
        DocumentVersion.objects.create(
            document=document,
            version_number=1,
            source_filename="instruction_v1.txt",
            file=self.make_file("instruction_v1.txt", "Редакция 1"),
            file_size=len("Редакция 1".encode("utf-8")),
            content_type="text/plain",
        )

        response = self.client.post(
            reverse("documents-versions", kwargs={"pk": document.id}),
            {
                "file": self.make_file("instruction_v2.txt", "Редакция 2"),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_201_CREATED)
        self.assertEqual(response.data["version_number"], 2)
        self.assertEqual(document.versions.count(), 2)

    def test_versions_list_can_be_requested_for_specific_document(self):
        first_document = Document.objects.create(title="Документ 1")
        second_document = Document.objects.create(title="Документ 2")

        DocumentVersion.objects.create(
            document=first_document,
            version_number=1,
            source_filename="doc1_v1.txt",
            file=self.make_file("doc1_v1.txt", "Первая версия"),
            file_size=len("Первая версия".encode("utf-8")),
            content_type="text/plain",
        )
        DocumentVersion.objects.create(
            document=second_document,
            version_number=1,
            source_filename="doc2_v1.txt",
            file=self.make_file("doc2_v1.txt", "Другая версия"),
            file_size=len("Другая версия".encode("utf-8")),
            content_type="text/plain",
        )

        response = self.client.get(
            reverse("versions-list"),
            {"document": first_document.id},
        )

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(response.data[0]["document"], first_document.id)
        self.assertEqual(response.data[0]["version_number"], 1)

    def test_document_list_contains_versions_count_and_latest_version_number(self):
        document = Document.objects.create(title="Положение")
        DocumentVersion.objects.create(
            document=document,
            version_number=1,
            source_filename="v1.txt",
            file=self.make_file("v1.txt", "Версия 1"),
            file_size=len("Версия 1".encode("utf-8")),
            content_type="text/plain",
        )
        DocumentVersion.objects.create(
            document=document,
            version_number=2,
            source_filename="v2.txt",
            file=self.make_file("v2.txt", "Версия 2"),
            file_size=len("Версия 2".encode("utf-8")),
            content_type="text/plain",
        )

        response = self.client.get(reverse("documents-list"))

        self.assertEqual(response.status_code, status.HTTP_200_OK)
        self.assertEqual(len(response.data), 1)
        self.assertEqual(response.data[0]["versions_count"], 2)
        self.assertEqual(response.data[0]["latest_version_number"], 2)

    def test_upload_rejects_unsupported_extension(self):
        document = Document.objects.create(title="Плохой файл")

        response = self.client.post(
            reverse("documents-versions", kwargs={"pk": document.id}),
            {
                "file": self.make_file(
                    "malware.exe",
                    "not allowed",
                    content_type="application/octet-stream",
                ),
            },
            format="multipart",
        )

        self.assertEqual(response.status_code, status.HTTP_400_BAD_REQUEST)
        self.assertIn("Unsupported file type", response.data["file"][0])
