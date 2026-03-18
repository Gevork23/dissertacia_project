from __future__ import annotations

from io import BytesIO

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from docx import Document as DocxDocument
from pypdf import PdfWriter

from .text_extractors import (
    EmptyExtractedTextError,
    extract_text_from_bytes,
    process_uploaded_file,
)
from .text_processing import normalize_text, sha256_hex


class TextExtractionUnitTests(TestCase):
    def make_docx_bytes(self) -> bytes:
        buffer = BytesIO()
        document = DocxDocument()
        document.add_paragraph("Приказ по МФЦ")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Срок"
        table.rows[0].cells[1].text = "5 дней"
        document.save(buffer)
        return buffer.getvalue()

    def make_pdf_bytes(self, text: str = "Hello PDF") -> bytes:
        safe_text = (
            text.replace("\\", "\\\\")
            .replace("(", "\\(")
            .replace(")", "\\)")
        )
        stream = (
            f"BT\n/F1 18 Tf\n50 100 Td\n({safe_text}) Tj\nET".encode("latin-1")
        )
        objects = [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            (
                b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 144] "
                b"/Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
            ),
            b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        ]

        result = bytearray(b"%PDF-1.4\n")
        offsets = [0]
        for index, obj in enumerate(objects, start=1):
            offsets.append(len(result))
            result.extend(f"{index} 0 obj\n".encode())
            result.extend(obj)
            result.extend(b"\nendobj\n")

        xref_offset = len(result)
        result.extend(f"xref\n0 {len(objects) + 1}\n".encode())
        result.extend(b"0000000000 65535 f \n")
        for offset in offsets[1:]:
            result.extend(f"{offset:010d} 00000 n \n".encode())
        result.extend(
            (
                f"trailer\n<< /Root 1 0 R /Size {len(objects) + 1} >>\n"
                f"startxref\n{xref_offset}\n%%EOF\n"
            ).encode()
        )
        return bytes(result)

    def test_extract_text_from_txt_uses_cp1251_fallback(self):
        raw = "Приказ № 1".encode("cp1251")

        extracted = extract_text_from_bytes(raw, "order.txt")

        self.assertEqual(extracted, "Приказ № 1")

    def test_extract_text_from_docx_reads_paragraphs_and_tables(self):
        extracted = extract_text_from_bytes(self.make_docx_bytes(), "order.docx")

        self.assertIn("Приказ по МФЦ", extracted)
        self.assertIn("Срок | 5 дней", extracted)

    def test_extract_text_from_pdf_without_ocr(self):
        extracted = extract_text_from_bytes(self.make_pdf_bytes(), "order.pdf")

        self.assertIn("Hello PDF", extracted)

    def test_process_uploaded_file_returns_normalized_text_and_hash(self):
        uploaded_file = SimpleUploadedFile(
            name="law.txt",
            content="Строка 1\r\n\r\nСтрока 2".encode("utf-8"),
            content_type="text/plain",
        )

        processed = process_uploaded_file(uploaded_file)

        self.assertEqual(processed.extracted_text, "Строка 1\n\nСтрока 2")
        self.assertEqual(processed.normalized_text, normalize_text(processed.extracted_text))
        self.assertEqual(processed.content_hash, sha256_hex(processed.normalized_text))

    def test_process_uploaded_file_rejects_pdf_without_extractable_text(self):
        buffer = BytesIO()
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=144)
        writer.write(buffer)

        uploaded_file = SimpleUploadedFile(
            name="scanned.pdf",
            content=buffer.getvalue(),
            content_type="application/pdf",
        )

        with self.assertRaises(EmptyExtractedTextError):
            process_uploaded_file(uploaded_file)
